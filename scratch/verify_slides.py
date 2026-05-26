import pptx

prs = pptx.Presentation("thesis_presentation.pptx")
print("=== VERIFICATION REPORT ===")
print(f"Total slides count: {len(prs.slides)}")

target_slides = [
    {"index": 3, "expected_title": "PROJECT CONTEXT"},
    {"index": 4, "expected_title": "THEORETICAL FOUNDATION"},
    {"index": 5, "expected_title": "RESEARCH OBJECTIVES"},
    {"index": 6, "expected_title": "AGILE DEVELOPMENT METHODOLOGY"},
    {"index": 7, "expected_title": "CONCEPTUAL FRAMEWORK"},
    {"index": 45, "expected_title": "RESULTS: FUNCTIONAL & BETA TESTING"},
    {"index": 46, "expected_title": "RESULTS: ALPHA & SPEED TESTING"},
    {"index": 47, "expected_title": "RESULTS: SYSTEM EVALUATION"}
]

errors = 0

for ts in target_slides:
    idx = ts["index"]
    slide = prs.slides[idx]
    
    # Check title
    title_tb = [s for s in slide.shapes if s.name == 'TextBox 7']
    if not title_tb:
        print(f"[ERROR] Slide {idx+1} does not have TextBox 7 (Title text box)!")
        errors += 1
        continue
    
    actual_title = title_tb[0].text_frame.text
    if actual_title != ts["expected_title"]:
        print(f"[ERROR] Slide {idx+1} title mismatch! Expected: '{ts['expected_title']}', Got: '{actual_title}'")
        errors += 1
    else:
        print(f"[OK] Slide {idx+1} Title: '{actual_title}'")
        
    # Check content text box (TextBox 6)
    content_tb = [s for s in slide.shapes if s.name == 'TextBox 6']
    if not content_tb:
        print(f"[ERROR] Slide {idx+1} does not have TextBox 6 (Content text box)!")
        errors += 1
        continue
        
    print(f"  Content length: {len(content_tb[0].text_frame.text)} chars")
    paragraphs = content_tb[0].text_frame.paragraphs
    print(f"  Paragraph count: {len(paragraphs)}")
    for i, p in enumerate(paragraphs[:2]):
        print(f"    P{i+1}: {p.text[:80]}...")
        # Check fonts of runs
        for run in p.runs[:2]:
            print(f"      Run font: {run.font.name}, size: {run.font.size}, bold: {run.font.bold}, color: {run.font.color.rgb if run.font.color else 'None'}")
            
    # Check background shapes (Group 2 and Group 4)
    g2 = [s for s in slide.shapes if s.name == 'Group 2']
    g4 = [s for s in slide.shapes if s.name == 'Group 4']
    if not g2 or not g4:
        print(f"[ERROR] Slide {idx+1} is missing background shapes (Group 2 or Group 4)!")
        errors += 1
    else:
        print(f"  [OK] Background shapes present.")

if errors == 0:
    print("\n=== ALL TESTS PASSED SUCCESSFULLY! ===")
else:
    print(f"\n=== VERIFICATION FAILED WITH {errors} ERRORS ===")
