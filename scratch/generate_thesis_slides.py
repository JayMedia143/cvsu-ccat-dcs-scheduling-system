import pptx
import copy
import shutil
from pptx.dml.color import RGBColor
from pptx.util import Pt

# 1. Copy the original presentation
src_ppt = "alpha_test11_presentation.pptx"
dest_ppt = "thesis_presentation.pptx"
shutil.copyfile(src_ppt, dest_ppt)
print(f"Copied {src_ppt} to {dest_ppt} successfully!")

# 2. Load presentation
prs = pptx.Presentation(dest_ppt)
print(f"Loaded presentation. Initial slide count: {len(prs.slides)}")

# 3. Define formatting helpers
def set_title(slide, title_text):
    tb = [s for s in slide.shapes if s.name == 'TextBox 7'][0]
    p = tb.text_frame.paragraphs[0]
    if p.runs:
        p.runs[0].text = title_text
        # Clear extra runs
        for r in list(p.runs[1:]):
            p.runs.remove(r)
    else:
        p.text = title_text

def add_bullet_point(text_frame, bullet_text, size_pt=20):
    if len(text_frame.paragraphs) == 1 and text_frame.paragraphs[0].text == "":
        p = text_frame.paragraphs[0]
    else:
        p = text_frame.add_paragraph()
    
    p.space_after = Pt(8)
    p.space_before = Pt(4)
    
    # Split by double asterisks for bold markdown highlighting
    parts = bullet_text.split("**")
    
    # Prepend bullet symbol to first run in paragraph
    bullet_symbol = "•  "
    
    for idx, part in enumerate(parts):
        run_text = bullet_symbol + part if idx == 0 else part
        if not run_text:
            continue
        run = p.add_run()
        run.text = run_text
        is_bold = (idx % 2 == 1)
        run.font.name = "Calibri"
        run.font.size = Pt(size_pt)
        run.font.bold = is_bold
        # Primary dark green/gray color: RGB (65, 75, 59)
        run.font.color.rgb = RGBColor(65, 75, 59)

# 4. Slide definitions
all_new_slides = [
    # Group A: Introduction & Objectives (5 slides)
    {
        "title": "PROJECT CONTEXT",
        "bullets": [
            "The **University Course Timetabling Problem (UCTP)** is a globally recognized, computationally complex **NP-hard** problem.",
            "As variables like courses and instructors increase, combinations grow exponentially, making brute-force methods mathematically impossible.",
            "At **CvSU-CCAT Campus**, the Department of Computer Studies (DCS) still relies on a traditional **manual scheduling process**.",
            "This manual trial-and-error process is slow and fragile, taking **a month or longer** of faculty effort and causing resource conflicts.",
            "There is an urgent need to transition to an automated, intelligent scheduling system tailored to the department's constraints."
        ]
    },
    {
        "title": "THEORETICAL FOUNDATION",
        "bullets": [
            "**Constraint Programming (CP)**: Formally models timetabling by defining variables, domains, and non-negotiable **Hard Constraints** vs. desirable **Soft Constraints**.",
            "**Genetic Algorithm (GA)**: The most popular meta-heuristic algorithm in scheduling optimization (representing **26%** of publications in literature).",
            "Mimics natural selection (**crossover**, **mutation**, and **survival of the fittest**) to explore large solution spaces and find optimal schedules.",
            "**Local Validation**: Recent Philippine studies (e.g., Carawana et al., 2025) achieved an **80% reduction** in manual workload, proving the viability of GAs in SUCs."
        ]
    },
    {
        "title": "RESEARCH OBJECTIVES",
        "bullets": [
            "**General Objective**: Develop an automated scheduling system to maximize room utilization by minimizing conflicts and optimizing available resources at CvSU-CCAT DCS.",
            "**Specific Objectives**:",
            "  1. **Design** the computer lab scheduling system utilizing a GA to satisfy hard and soft constraints.",
            "  2. **Develop** the system using the **Python** programming language and the **SQLite** database.",
            "  3. **Test** the system in terms of **Functionality**, **Alpha** (optimizations), **Speed**, and **Beta** (user feedback).",
            "  4. **Evaluate** software and data quality using the adapted **ISO/IEC 25010** software evaluation instrument.",
            "  5. **Prepare** a structured **Implementation Plan** for institutional deployment."
        ]
    },
    {
        "title": "AGILE DEVELOPMENT METHODOLOGY",
        "bullets": [
            "Adopts an **Agile Development Model** to manage the algorithmic complexities and allow iterative enhancements.",
            "**Planning Phase**: Analyzed manual workflows, gathered constraints, and established user priorities.",
            "**Design Phase**: Drafted the database schema in SQLite and mapped scheduling rules into algorithmic parameters.",
            "**Development Phase**: Implemented the Genetic Algorithm core in Python and designed the interactive drag-and-drop workspace.",
            "**Testing Phase**: Evaluated the system continuously to identify logic errors and ensure solid constraint enforcement.",
            "**Review & Deployment**: Refined UI based on adviser review and prepared the system for final user evaluation."
        ]
    },
    {
        "title": "CONCEPTUAL FRAMEWORK",
        "bullets": [
            "**Input-Process-Output (IPO) Model**:",
            "  - **Input**: Genetic Algorithm & Constraint Programming principles, Python, SQLite, and academic datasets.",
            "  - **Process**: Requirement analysis, software design, development, and multi-tier testing.",
            "  - **Output**: Fully functional **Automated Scheduling System** refined via an continuous evaluation loop.",
            "**Operational Feasibility**:",
            "  - Developed using **zero-cost** open-source software and student development effort.",
            "  - Runs on existing hardware (1 PC, 2 Laptops, asset value **PhP 40,000.00**), requiring **no extra institutional expenses**."
        ]
    },
    # Group B: Results & Discussion (3 slides)
    {
        "title": "RESULTS: FUNCTIONAL & BETA TESTING",
        "bullets": [
            "**Functional Testing**: Checked 19 distinct modules across **171 test cases** and **1,710 executions**.",
            "  - Achieved a **100% success rate** with zero critical failures.",
            "  - Successfully validated core modules including Security, GA Solver, Drag-and-Drop Grid, and Proposal Hub.",
            "**Beta Testing**: Deployed prototype to department heads and schedulers from various departments.",
            "  - Audited critical paths: login lockouts, irregular pathfinder, manual edits, and portal access.",
            "  - All reported items (e.g., mobile visibility, 5-unit course split) were successfully **resolved**, achieving a **usability rating of 4/5**."
        ]
    },
    {
        "title": "RESULTS: ALPHA & SPEED TESTING",
        "bullets": [
            "**Alpha Testing (GA Optimizations)**:",
            "  - Base GA was unstable; integrated **Memory Mapping (Bitmask Caching)**, **Targeted Repair**, and **Dynamic Slot Injection** to guide mutations.",
            "  - Drastically reduced solve time, successfully generating a flawless schedule in **2.1 minutes**.",
            "**Speed Testing Across Hardware**:",
            "  - Evaluated across different CPUs (Intel i7, i5, i3).",
            "  - On a mid-range PC, standard schedule generation took only **22 seconds**.",
            "  - Even on an entry-level laptop, schedules were generated in under **3 minutes**, proving **high portability and low hardware requirements**."
        ]
    },
    {
        "title": "RESULTS: SYSTEM EVALUATION",
        "bullets": [
            "Evaluated by **20 respondents** (10 IT Experts and 10 Academic Schedulers) across multiple university departments.",
            "**Excellent Ratings (Overall Mean: 4.66)**:",
            "  - *Functional Suitability*: **4.82 (Excellent)** — Complete, correct, and appropriate.",
            "  - *Interaction Capability*: **4.75 (Excellent)** — Highly learnable and engaging interface.",
            "  - *Flexibility & Maintainability*: **4.71 & 4.67 (Excellent)** — Highly modular, adaptive, and scalable.",
            "  - *Security & Reliability*: **4.66 & 4.53 (Excellent)** — High confidentiality and fault tolerance.",
            "  - *Compatibility*: **4.48 (Very Good)** — Smooth co-existence with existing campus tools."
        ]
    }
]

# 5. Generate new slides appended at the end of the presentation
# This avoids index shifts during generation!
template_slide = prs.slides[41] # Slide 42 is always at index 41 initially

for slide_info in all_new_slides:
    blank_layout = prs.slide_layouts[6]
    new_slide = prs.slides.add_slide(blank_layout)
    
    # Copy shapes from template_slide to new_slide
    for shape in template_slide.shapes:
        el = shape.element
        new_el = copy.deepcopy(el)
        new_slide.shapes._spTree.append(new_el)
        
    # Delete Freeform 8 if it exists in the new slide
    for shape in list(new_slide.shapes):
        if shape.name == 'Freeform 8':
            shape_el = shape.element
            shape_el.getparent().remove(shape_el)
            
    # Modify TextBox 7 (Title)
    set_title(new_slide, slide_info["title"])
    
    # Modify TextBox 6 (Content)
    tb_content = [s for s in new_slide.shapes if s.name == 'TextBox 6'][0]
    tf = tb_content.text_frame
    tf.clear()
    
    for bullet in slide_info["bullets"]:
        add_bullet_point(tf, bullet, size_pt=20)
        
    print(f"Generated slide: '{slide_info['title']}' at the end of presentation.")

# 6. Reorder slides cleanly using sldIdLst
print("Reordering slides...")
sldIdLst = prs.slides._sldIdLst
new_slide_elements = list(sldIdLst[-8:])

# Remove the 8 elements from the end
for elem in new_slide_elements:
    sldIdLst.remove(elem)

# Insert Group A elements (first 5 elements) at indices 3, 4, 5, 6, 7
for idx, elem in enumerate(new_slide_elements[:5]):
    sldIdLst.insert(3 + idx, elem)
    print(f"Moved Group A slide index {idx} to index {3 + idx}")

# Insert Group B elements (last 3 elements) at indices 45, 46, 47
for idx, elem in enumerate(new_slide_elements[5:]):
    sldIdLst.insert(45 + idx, elem)
    print(f"Moved Group B slide index {idx} to index {45 + idx}")

# 7. Save final presentation
prs.save(dest_ppt)
print(f"Saved final updated presentation to {dest_ppt} with slide count: {len(prs.slides)}")
