import openpyxl

def peek_template(path):
    print(f"\n--- Peeking at {path} ---")
    try:
        wb = openpyxl.load_workbook(path, data_only=True)
        ws = wb.active
        merged_ranges = ws.merged_cells.ranges
        
        def get_merged_range(c_idx, r_idx):
            for m_range in merged_ranges:
                if r_idx >= m_range.min_row and r_idx <= m_range.max_row and \
                   c_idx >= m_range.min_col and c_idx <= m_range.max_col:
                    return f"{m_range.coord}"
            return f"{openpyxl.utils.get_column_letter(c_idx)}{r_idx}"

        max_r = ws.max_row
        # Scan Top 40 and Bottom 30
        ranges = [(1, 41), (max(1, max_r - 30), max_r + 1)]
        
        seen_rows = set()
        for start, end in ranges:
            for r in range(start, end):
                if r in seen_rows: continue
                row_data = []
                for c in range(1, 15): # Scan up to column N
                    cell = ws.cell(row=r, column=c)
                    val = cell.value if cell.value is not None else ""
                    if val and not isinstance(val, (int, float)): # skip numerical data
                         coord = get_merged_range(c, r)
                         row_data.append(f"'{str(val)[:30]}' @ [{coord}]")
                if row_data:
                    print(f"Row {r}: " + " | ".join(row_data))
                    seen_rows.add(r)
    except Exception as e:
        print(f"Error reading {path}: {e}")

peek_template("static/assets/section_template.xlsx")
peek_template("static/assets/faculty_template.xlsx")
peek_template("static/assets/room_template.xlsx")
peek_template("static/assets/course_template.xlsx")
